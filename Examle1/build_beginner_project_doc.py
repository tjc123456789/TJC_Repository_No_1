from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Examle1工程说明文档-初学者详细版.docx"
FONT = "Microsoft YaHei"


def set_font(run, name=FONT, size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color="DADCE0", size="6"):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = tbl_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def table_width(table, width_dxa=9360, widths=None):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    table.autofit = False
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width / 1440)


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, "2E74B5", 18, 10
    elif level == 2:
        size, color, before, after = 13, "2E74B5", 14, 7
    else:
        size, color, before, after = 12, "1F4D78", 10, 5
    r = p.add_run(text)
    set_font(r, size=size, bold=True, color=color)
    spacing(p, before=before, after=after)
    return p


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    spacing(p)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r)
    spacing(p, after=4)
    return p


def callout(doc, title, body, fill="F4F6F9", color="0B2545"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width(t, 9360, [9360])
    borders(t, color="DADCE0", size="4")
    c = t.cell(0, 0)
    shade(c, fill)
    cell_margins(c, top=120, start=160, bottom=120, end=160)
    p = c.paragraphs[0]
    r = p.add_run(title + "：")
    set_font(r, bold=True, color=color)
    r = p.add_run(body)
    set_font(r)
    spacing(p, after=0)
    doc.add_paragraph()


def code(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width(t, 9360, [9360])
    borders(t, color="DADCE0", size="4")
    c = t.cell(0, 0)
    shade(c, "F6F8FA")
    cell_margins(c, top=110, start=150, bottom=110, end=150)
    p = c.paragraphs[0]
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        set_font(r, name="Consolas", size=9)
    spacing(p, after=0, line=1.0)
    doc.add_paragraph()


def data_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width(t, 9360, widths)
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, "E8EEF5")
        cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        r = p.add_run(h)
        set_font(r, bold=True, color="0B2545")
        spacing(p, after=0)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            r = p.add_run(value)
            set_font(r, size=10.5)
            spacing(p, after=0, line=1.18)
    doc.add_paragraph()
    return t


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Examle1 工程说明文档")
set_font(r, size=22, bold=True, color="0B2545")
spacing(title, after=4)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("初学者详细版：专业词汇解释 + 每个文件代码说明")
set_font(r, size=12, color="555555")
spacing(sub, after=18)

callout(doc, "阅读建议", "如果你是第一次接触 Java Web 项目，可以先看“一、先用生活例子理解这个项目”，再看“术语表”，最后再逐个看文件。不要一开始就死盯代码，先理解每个文件负责什么。")

heading(doc, "一、先用生活例子理解这个项目", 1)
para(doc, "这个工程可以想象成一个很小的“账号系统”。用户在网页上输入账号和密码，网页把这些信息交给后端程序，后端程序再去数据库里查询或保存数据。")
data_table(doc, ["现实类比", "在本工程中对应什么"], [
    ("用户来到办事窗口", "用户打开 login.html 或 register.html 页面。"),
    ("填写表格", "用户在页面输入账号、密码、确认密码。"),
    ("工作人员检查资料", "HelloController 检查账号是否为空、是否存在、密码是否正确。"),
    ("档案柜", "MySQL 数据库，用来长期保存用户数据。"),
    ("查档案和存档案的人", "UserRepository，负责查询和保存用户。"),
    ("每份档案的格式", "User 实体类，规定一个用户有哪些字段。"),
], [2500, 6860])

heading(doc, "二、这个工程最终实现了什么", 1)
bullet(doc, "提供一个登录页面：用户输入账号和密码后，可以请求后端验证。")
bullet(doc, "提供一个注册页面：用户输入账号、密码、确认密码后，可以创建新账号。")
bullet(doc, "连接 MySQL 数据库：把注册的用户保存下来，登录时再查询出来。")
bullet(doc, "使用 Spring Boot 启动 Web 服务：浏览器可以访问页面，前端也可以调用后端接口。")
bullet(doc, "使用 JPA 自动操作数据库：代码里不需要手写完整 SQL，也能完成保存和查询。")

heading(doc, "三、整体框图", 1)
code(doc, """浏览器中的页面
    |
    | 1. 用户输入账号和密码
    v
login.html / register.html
    |
    | 2. JavaScript 使用 fetch 发送请求
    v
HelloController
    |
    | 3. Controller 判断登录或注册逻辑
    v
UserRepository
    |
    | 4. Repository 通过 JPA/Hibernate 操作数据库
    v
MySQL 数据库 demo.user 表""")

heading(doc, "四、初学者必须先懂的专业词汇", 1)
terms = [
    ("工程 / 项目", "一组为了完成某个软件功能而放在一起的文件。这里的工程就是一个登录注册网站。"),
    ("前端", "用户能直接看到和操作的部分，比如 HTML 页面、按钮、输入框。"),
    ("后端", "用户看不到但负责处理业务逻辑的程序，比如校验密码、访问数据库。"),
    ("数据库", "专门保存数据的软件。这个工程使用 MySQL 保存账号和密码。"),
    ("MySQL", "一种常见的关系型数据库，可以理解为很多张表组成的数据仓库。"),
    ("表", "数据库里保存同一类数据的结构。例如 user 表保存所有用户。"),
    ("字段", "表中的一列。例如用户表里的 id、username、password。"),
    ("主键", "一条数据的唯一编号。这里 User 的 id 就是主键。"),
    ("Java", "本工程后端使用的编程语言。"),
    ("Spring Boot", "一个 Java Web 开发框架，能帮我们快速启动网站、处理请求、连接数据库。"),
    ("框架", "别人提前写好的基础代码和规则。我们按它的方式写代码，可以少做很多重复工作。"),
    ("Maven", "Java 项目的构建和依赖管理工具。它负责下载依赖、编译代码、打包项目。"),
    ("依赖", "项目需要用到的外部功能包。例如 Spring Web、JPA、MySQL 驱动。"),
    ("JAR", "Java 打包后的文件格式。target 里的 jar 可以用来运行整个项目。"),
    ("Tomcat", "Web 服务器。Spring Boot 内置 Tomcat，所以运行项目后浏览器才能访问。"),
    ("Controller", "控制器，负责接收浏览器请求并返回结果。HelloController 就是控制器。"),
    ("Repository", "数据仓库层，负责和数据库打交道。UserRepository 就是这个角色。"),
    ("Entity", "实体类，用 Java 类表示数据库中的表。User 类就是实体类。"),
    ("JPA", "Java 操作数据库的一套标准。它让我们用对象方式操作表数据。"),
    ("Hibernate", "JPA 的常用实现。真正帮我们把 Java 对象转换为 SQL 操作数据库。"),
    ("注解", "Java 中以 @ 开头的标记，例如 @Entity、@RestController。它们告诉框架这个类或方法有什么特殊用途。"),
    ("接口", "在这里通常指后端提供给前端调用的地址，比如 /login 和 /register。"),
    ("HTTP", "浏览器和服务器之间通信的规则。访问网页、提交表单都依赖 HTTP。"),
    ("POST", "HTTP 请求方式之一，常用于提交数据，例如登录和注册。"),
    ("JSON", "一种常见的数据格式，长得像 {\"username\":\"abc\"}，前后端经常用它传数据。"),
    ("静态资源", "不需要后端动态生成的文件，比如 HTML、CSS、JavaScript、图片。"),
    ("编译", "把 Java 源代码转换成 JVM 能执行的 .class 字节码。"),
    ("字节码", "Java 编译后的中间文件，扩展名通常是 .class。"),
    ("target 目录", "Maven 构建后生成的目录，里面放编译结果和打包结果。"),
    ("IDE", "写代码的软件，例如 IntelliJ IDEA。项目里的 .idea 是 IDE 配置。"),
]
data_table(doc, ["术语", "初学者解释"], terms, [2200, 7160])

heading(doc, "五、文件结构说明", 1)
data_table(doc, ["文件或目录", "它是做什么的", "初学者怎么理解"], [
    ("src/main/java/com/example/demo", "后端 Java 源码目录。", "真正处理登录、注册、数据库访问的代码主要在这里。"),
    ("DemoApplication.java", "项目启动入口。", "像电源开关，运行它整个网站才会启动。"),
    ("HelloController.java", "后端接口控制器。", "像办事窗口，接待登录和注册请求。"),
    ("User.java", "用户实体类。", "规定一个用户对象长什么样，有 id、username、password。"),
    ("UserRepository.java", "数据库访问接口。", "像查档案的人，负责按用户名查用户、保存用户。"),
    ("src/main/resources", "配置和资源目录。", "放项目运行需要的配置文件和网页。"),
    ("application.properties", "Spring Boot 配置文件。", "告诉项目数据库地址、账号、密码、JPA 设置。"),
    ("static/index.html", "首页。", "打开网站根路径时自动跳到登录页。"),
    ("static/login.html", "登录页面。", "用户输入账号密码，页面调用 /login。"),
    ("static/register.html", "注册页面。", "用户创建账号，页面调用 /register。"),
    ("src/test/java", "测试代码目录。", "用来检查项目是否能正常启动。"),
    ("pom.xml", "Maven 配置文件。", "告诉 Maven 这个项目需要哪些工具包、怎么打包。"),
    ("target", "构建产物目录。", "编译后自动生成，不是主要源码。"),
], [2600, 3400, 3360])

heading(doc, "六、后端代码逐文件解释", 1)
heading(doc, "1. DemoApplication.java：项目启动入口", 2)
para(doc, "这个文件很短，但非常重要。它负责启动整个 Spring Boot 应用。")
code(doc, """@SpringBootApplication
public class DemoApplication {
    public static void main(String[] args) {
        SpringApplication.run(DemoApplication.class, args);
    }
}""")
data_table(doc, ["代码", "解释"], [
    ("@SpringBootApplication", "告诉 Spring Boot：这是一个应用的启动类。它会自动扫描当前包下的 Controller、Repository、Entity 等组件。"),
    ("public class DemoApplication", "定义一个 Java 类，类名叫 DemoApplication。"),
    ("main(String[] args)", "Java 程序的入口方法。运行项目时，会先从 main 方法开始执行。"),
    ("SpringApplication.run(...)", "真正启动 Spring Boot，包括内置 Tomcat、配置文件读取、组件扫描等。"),
], [3200, 6160])
callout(doc, "一句话理解", "DemoApplication.java 就是这个网站后端程序的启动按钮。")

heading(doc, "2. HelloController.java：处理登录和注册", 2)
para(doc, "这个文件是业务逻辑最集中的地方。前端页面提交登录或注册请求后，最终都会来到这里。")
code(doc, """@RestController
public class HelloController {

    @Autowired
    private UserRepository userRepository;

    @PostMapping("/login")
    public Map<String, Object> login(@RequestParam String username,
                                     @RequestParam String password) {
        ...
    }

    @PostMapping("/register")
    public Map<String, Object> register(@RequestBody Map<String, String> params) {
        ...
    }
}""")
data_table(doc, ["代码或词汇", "初学者解释"], [
    ("@RestController", "表示这个类是后端接口类。方法返回的 Map 会自动变成 JSON 返回给前端。"),
    ("@Autowired", "自动把 UserRepository 对象交给这个类使用，初学阶段可以理解为“自动帮你准备好工具”。"),
    ("UserRepository userRepository", "数据库操作工具。需要查用户、保存用户时就调用它。"),
    ("@PostMapping(\"/login\")", "表示这个方法处理 POST /login 请求。"),
    ("@RequestParam", "表示参数来自表单提交，例如 username=tom&password=123。"),
    ("Map<String, Object>", "返回给前端的数据容器。里面可以放 code、message、username 等键值对。"),
    ("@RequestBody", "表示参数来自请求体，通常是 JSON 数据。register.html 注册时就是这样提交的。"),
], [3100, 6260])
para(doc, "登录逻辑可以拆成四步：")
bullet(doc, "第一步：检查账号和密码是否为空。为空就直接返回错误。")
bullet(doc, "第二步：按用户名去数据库查询用户。")
bullet(doc, "第三步：如果用户不存在，提示先注册。")
bullet(doc, "第四步：如果用户存在，再比较密码；密码正确就登录成功。")
code(doc, """User user = userRepository.findByUsername(username.trim());
if (user == null) {
    result.put("code", 2);
    result.put("message", "账号不存在，请先注册");
    return result;
}

if (!user.getPassword().equals(password)) {
    result.put("code", 1);
    result.put("message", "密码错误，请重新输入");
    return result;
}""")
callout(doc, "重要提醒", "这里用 user.getPassword().equals(password) 直接比较明文密码。学习时容易理解，但真实系统不能这样保存密码，应使用加密哈希。")
para(doc, "注册逻辑可以拆成三步：")
bullet(doc, "第一步：从前端传来的 JSON 中取出 username 和 password。")
bullet(doc, "第二步：检查用户名是否已经存在。")
bullet(doc, "第三步：创建 User 对象并调用 save 保存到数据库。")

heading(doc, "3. User.java：用户这张表的 Java 表达", 2)
para(doc, "User.java 是实体类。实体类的作用是：用 Java 类描述数据库表中的一条数据。")
code(doc, """@Entity
public class User {
    @Id
    @GeneratedValue(strategy = GenerationType.IDENTITY)
    private Long id;
    private String username;
    private String password;
}""")
data_table(doc, ["代码或词汇", "初学者解释"], [
    ("@Entity", "告诉 JPA：这个类对应数据库中的一张表。"),
    ("@Id", "表示 id 是主键，也就是每条用户记录的唯一编号。"),
    ("@GeneratedValue", "表示 id 自动生成，不需要用户手动填写。"),
    ("GenerationType.IDENTITY", "表示使用数据库自己的自增方式生成 id。"),
    ("private Long id", "用户编号。Long 是 Java 的整数类型之一。"),
    ("private String username", "账号字段。String 表示文本。"),
    ("private String password", "密码字段。当前项目中是明文保存。"),
    ("无参构造方法", "JPA 从数据库读取数据后，需要用它创建对象。"),
    ("getter/setter", "读取和修改字段的方法，例如 getUsername、setPassword。"),
], [3000, 6360])
callout(doc, "一句话理解", "User.java 就像一张“用户档案模板”，规定每个用户记录有哪些信息。")

heading(doc, "4. UserRepository.java：帮我们操作数据库", 2)
code(doc, """@Repository
public interface UserRepository extends JpaRepository<User, Long> {
    User findByUsername(String username);
}""")
data_table(doc, ["代码或词汇", "初学者解释"], [
    ("@Repository", "表示这是数据库访问层组件。"),
    ("interface", "接口。这里不写具体实现，Spring Data JPA 会自动帮我们生成实现。"),
    ("extends JpaRepository<User, Long>", "继承 JPA 提供的通用数据库操作能力。User 表示操作 User 实体，Long 表示主键 id 的类型。"),
    ("save(user)", "虽然代码里没有写，但继承 JpaRepository 后自动拥有，用来保存用户。"),
    ("findById(id)", "自动拥有的方法，用 id 查询用户。"),
    ("findByUsername", "根据方法名自动生成查询逻辑，意思是按 username 字段查用户。"),
], [3300, 6060])
callout(doc, "神奇但常见的地方", "findByUsername 没有写 SQL，却能查询数据库，这是 Spring Data JPA 根据方法名自动推导查询。")

heading(doc, "七、前端页面逐文件解释", 1)
heading(doc, "1. index.html：首页跳转", 2)
code(doc, """<meta http-equiv="refresh" content="0;url=login.html">""")
para(doc, "这行代码表示：打开 index.html 后，浏览器立刻跳转到 login.html。content=\"0;url=login.html\" 中的 0 表示等待 0 秒。")

heading(doc, "2. login.html：登录页面", 2)
para(doc, "login.html 同时包含三类内容：HTML 负责页面结构，CSS 负责页面样式，JavaScript 负责交互逻辑。")
data_table(doc, ["部分", "作用", "例子"], [
    ("HTML", "放输入框、按钮、提示区域。", "input、button、div。"),
    ("CSS", "控制页面长什么样。", "背景颜色、按钮颜色、卡片阴影。"),
    ("JavaScript", "让页面能和后端通信。", "fetch('/login', ...)。"),
], [1700, 3800, 3860])
code(doc, """fetch('/login', {
    method: 'POST',
    headers: { 'Content-Type': 'application/x-www-form-urlencoded' },
    body: 'username=' + encodeURIComponent(username) +
          '&password=' + encodeURIComponent(password)
})""")
data_table(doc, ["代码或词汇", "初学者解释"], [
    ("fetch", "浏览器提供的函数，用来向后端发送请求。"),
    ("'/login'", "请求地址，对应 HelloController 里的 @PostMapping(\"/login\")。"),
    ("method: 'POST'", "使用 POST 方式提交数据。"),
    ("Content-Type", "告诉后端这次提交的数据是什么格式。"),
    ("application/x-www-form-urlencoded", "传统表单格式，长得像 username=tom&password=123。"),
    ("encodeURIComponent", "把特殊字符转换成安全格式，避免账号密码里有特殊符号导致请求出错。"),
    ("data.code === 0", "后端返回 code=0 时，前端认为登录成功。"),
], [3300, 6060])

heading(doc, "3. register.html：注册页面", 2)
para(doc, "register.html 的主要工作是：让用户输入账号、密码、确认密码；先在前端检查两次密码是否一致；再把账号和密码以 JSON 格式发给后端。")
code(doc, """const res = await fetch('/register', {
    method: 'POST',
    headers: { 'Content-Type': 'application/json' },
    body: JSON.stringify({ username, password })
});""")
data_table(doc, ["代码或词汇", "初学者解释"], [
    ("async / await", "让异步请求写起来更像顺序执行。可以理解为“等后端返回结果后再继续”。"),
    ("'/register'", "请求地址，对应 HelloController 里的 @PostMapping(\"/register\")。"),
    ("application/json", "告诉后端：这次提交的数据是 JSON 格式。"),
    ("JSON.stringify", "把 JavaScript 对象转换成 JSON 字符串。"),
    ("data.success", "后端返回 success=true 时，前端认为注册成功。"),
    ("setTimeout", "延迟执行。这里用于注册成功后 3 秒跳转回登录页。"),
], [3300, 6060])

heading(doc, "八、配置文件解释", 1)
heading(doc, "application.properties", 2)
para(doc, "这个文件是 Spring Boot 的配置中心。很多“项目运行时需要知道的信息”都写在这里，比如数据库地址、账号、密码、JPA 行为。")
data_table(doc, ["配置", "初学者解释"], [
    ("spring.application.name=demo", "给应用取名叫 demo。"),
    ("spring.datasource.url=jdbc:mysql://localhost:3306/demo", "连接本机 MySQL 的 demo 数据库。localhost 表示本机，3306 是 MySQL 默认端口。"),
    ("useSSL=false", "关闭 SSL 连接。学习环境常见写法。"),
    ("characterEncoding=UTF-8", "使用 UTF-8 编码，避免中文乱码。"),
    ("serverTimezone=Asia/Shanghai", "设置数据库连接使用上海时区。"),
    ("spring.datasource.username=root", "数据库登录用户名。"),
    ("spring.datasource.password=000906", "数据库登录密码。"),
    ("driver-class-name=com.mysql.cj.jdbc.Driver", "指定 MySQL 驱动。驱动可以理解为 Java 和 MySQL 之间的翻译器。"),
    ("spring.jpa.hibernate.ddl-auto=update", "让 Hibernate 根据 User 类自动更新数据库表结构。"),
    ("spring.jpa.show-sql=true", "把执行的 SQL 打印出来，方便学习和调试。"),
    ("hibernate.dialect=MySQL5Dialect", "告诉 Hibernate 当前数据库接近 MySQL5 的语法。"),
], [3800, 5560])

heading(doc, "九、pom.xml 解释", 1)
para(doc, "pom.xml 是 Maven 项目的核心配置。它告诉 Maven：这个项目叫什么、使用哪个 Spring Boot 版本、需要哪些依赖、如何打包。")
data_table(doc, ["pom.xml 中的内容", "初学者解释"], [
    ("spring-boot-starter-parent", "Spring Boot 的父配置，提供很多默认版本和构建规则。"),
    ("groupId=com.example", "项目组织名，类似包名前缀。"),
    ("artifactId=demo", "项目产物名，最终 jar 名称会用到它。"),
    ("version=0.0.1-SNAPSHOT", "项目版本号。SNAPSHOT 表示开发中的快照版本。"),
    ("java.version=1.8", "项目代码按 Java 8 版本编译。"),
    ("spring-boot-starter-web", "让项目具备 Web 能力，比如 Controller、接口、内置 Tomcat。"),
    ("spring-boot-starter-data-jpa", "让项目具备用 JPA 操作数据库的能力。"),
    ("mysql-connector-java", "MySQL 驱动，让 Java 能连接 MySQL。"),
    ("spring-boot-starter-test", "测试相关依赖。"),
    ("spring-boot-maven-plugin", "把项目打包成可运行的 Spring Boot jar。"),
], [3600, 5760])

heading(doc, "十、测试和构建产物解释", 1)
heading(doc, "DemoApplicationTests.java", 2)
para(doc, "这是 Spring Boot 默认生成的测试类。它没有真正测试登录注册功能，只是测试 Spring 容器能不能正常启动。")
data_table(doc, ["代码", "初学者解释"], [
    ("@SpringBootTest", "测试时启动 Spring Boot 应用上下文。"),
    ("@Test", "表示下面的方法是一个测试方法。"),
    ("contextLoads()", "方法体为空，但如果 Spring 启动失败，这个测试就会失败。"),
], [3000, 6360])

heading(doc, "target 目录", 2)
para(doc, "target 目录是 Maven 构建后自动生成的，不是手写源码。通常不用手动修改它，删掉后重新构建还会生成。")
data_table(doc, ["文件", "初学者解释"], [
    ("target/classes/*.class", "Java 源码编译后的字节码文件。"),
    ("target/classes/static/*.html", "静态页面复制到运行目录后的副本。"),
    ("target/demo-0.0.1-SNAPSHOT.jar", "最终可运行的项目包。"),
    ("target/demo-0.0.1-SNAPSHOT.jar.original", "Spring Boot 重新打包前的原始 jar。"),
    ("target/maven-status", "Maven 记录编译过程的状态文件。"),
], [3500, 5860])

heading(doc, "十一、一次完整请求是怎样走的", 1)
code(doc, """以登录为例：

1. 用户在 login.html 输入账号和密码
2. 点击“登录”按钮
3. login.html 中的 JavaScript 执行 login() 函数
4. fetch('/login') 把账号密码发给后端
5. HelloController.login() 接收到请求
6. login() 调用 userRepository.findByUsername()
7. UserRepository 通过 JPA/Hibernate 查询 MySQL
8. 查询结果变成 User 对象返回给 Controller
9. Controller 判断密码是否正确
10. Controller 返回 JSON 给前端
11. 前端根据 code 显示成功或失败提示""")

heading(doc, "十二、这个项目目前的不足", 1)
bullet(doc, "密码明文保存，不安全。真实项目应使用 BCrypt 等哈希算法。")
bullet(doc, "没有登录状态保持。刷新页面后不会记住用户是否登录。")
bullet(doc, "没有权限控制。任何人都可以访问页面和接口。")
bullet(doc, "用户名没有唯一索引保护。代码检查了重名，但数据库层面最好也加唯一约束。")
bullet(doc, "测试较少。现在只有启动测试，没有专门测试注册和登录接口。")
bullet(doc, "前端页面直接写在 HTML 文件里，适合小项目学习；复杂项目通常会拆分成更多组件。")

heading(doc, "十三、推荐学习顺序", 1)
data_table(doc, ["顺序", "先学什么", "为什么"], [
    ("1", "HTML/CSS/JavaScript 基础", "先理解页面上的输入框、按钮和 fetch 请求。"),
    ("2", "HTTP、GET、POST、JSON", "理解前端和后端如何传数据。"),
    ("3", "Java 类、对象、方法、注解", "读懂 DemoApplication、User、Controller。"),
    ("4", "Spring Boot Controller", "理解 /login 和 /register 是怎么接收请求的。"),
    ("5", "MySQL 表和字段", "理解 User 对应数据库表。"),
    ("6", "JPA Repository", "理解为什么不用手写 SQL 也能查数据库。"),
    ("7", "Maven 和 pom.xml", "理解项目依赖和打包。"),
], [900, 3000, 5460])

heading(doc, "十四、速查术语表", 1)
quick_terms = [
    ("@RestController", "后端接口类标记，返回 JSON。"),
    ("@PostMapping", "把某个 POST 请求地址绑定到一个 Java 方法。"),
    ("@RequestParam", "从表单参数中取值。"),
    ("@RequestBody", "从请求体 JSON 中取值。"),
    ("@Entity", "把 Java 类映射成数据库表。"),
    ("@Id", "标记主键字段。"),
    ("@GeneratedValue", "标记主键自动生成。"),
    ("JpaRepository", "JPA 提供的通用数据库操作接口。"),
    ("findByUsername", "按 username 字段查询用户。"),
    ("save", "保存数据到数据库。"),
    ("fetch", "前端向后端发请求的函数。"),
    ("JSON.stringify", "把 JavaScript 对象转成 JSON 字符串。"),
    ("localhost", "本机地址。"),
    ("8080", "Spring Boot Web 服务常用端口。"),
    ("3306", "MySQL 默认端口。"),
]
data_table(doc, ["术语", "一句话解释"], quick_terms, [2600, 6760])

heading(doc, "十五、结论", 1)
para(doc, "这个工程虽然简单，但已经包含了 Web 开发的主线：页面、请求、后端接口、数据库、配置、构建和测试。对于初学者来说，它是一个很适合入门的项目，因为你可以从用户点击按钮开始，一路追踪到数据库保存和查询。")
para(doc, "学懂这个项目后，再去学习更复杂的用户系统、权限管理、前后端分离、密码加密和接口测试，会顺很多。")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Examle1 工程说明文档 - 初学者详细版")
    set_font(r, size=9, color="777777")

doc.save(OUT)
print(OUT)
