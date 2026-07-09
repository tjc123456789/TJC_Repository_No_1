from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Examle1工程说明文档.docx"


def set_run_font(run, name="Microsoft YaHei", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADCE0", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360, col_widths=None):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    table.autofit = False
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width / 1440)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    if level == 1:
        size, color, before, after = 16, "2E74B5", 18, 10
    elif level == 2:
        size, color, before, after = 13, "2E74B5", 14, 7
    else:
        size, color, before, after = 12, "1F4D78", 10, 5
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)
    set_paragraph_spacing(paragraph, before=before, after=after, line=1.25)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    set_paragraph_spacing(paragraph)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run)
    set_paragraph_spacing(paragraph, after=4)
    return paragraph


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 9360, [9360])
    set_table_borders(table, color="DADCE0", size="4")
    cell = table.cell(0, 0)
    shade_cell(cell, "F6F8FA")
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for idx, line in enumerate(text.splitlines()):
        if idx:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=9)
    set_paragraph_spacing(paragraph, after=0, line=1.0)
    return table


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 9360, widths)
    set_table_borders(table)
    header_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        shade_cell(header_cells[i], "E8EEF5")
        set_cell_margins(header_cells[i])
        header_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = header_cells[i].paragraphs[0]
        run = p.add_run(title)
        set_run_font(run, bold=True, color="0B2545")
        set_paragraph_spacing(p, after=0)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            run = p.add_run(value)
            set_run_font(run)
            set_paragraph_spacing(p, after=0)
    doc.add_paragraph()
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Examle1 工程说明文档")
set_run_font(run, size=22, bold=True, color="0B2545")
set_paragraph_spacing(title, before=0, after=8, line=1.15)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Spring Boot + MySQL 用户注册登录示例项目")
set_run_font(run, size=12, color="555555")
set_paragraph_spacing(subtitle, after=18)

add_heading(doc, "一、工程实现了什么", 1)
add_body(doc, "这个工程是一个 Spring Boot 用户注册/登录示例系统。前端由静态 HTML 页面组成，后端提供 /register 和 /login 两个接口，用户数据通过 Spring Data JPA 保存到本地 MySQL 数据库 demo 中的 user 表。")
add_body(doc, "用户打开网站后会进入登录页；没有账号可以跳到注册页；注册成功后写入数据库；登录时后端根据用户名查询用户并校验密码。")
add_body(doc, "注意：当前密码是明文保存的，适合课程演示或学习示例，不适合真实生产环境。真实项目应使用 BCrypt 等方式加密密码，并增加会话、鉴权和输入安全校验。")

add_heading(doc, "二、系统结构框图", 1)
add_code_block(doc, """用户浏览器
   |
   v
index.html 首页重定向
   |
   v
login.html 登录页面 <----------> register.html 注册页面
   |                                  |
   | POST /login                      | POST /register
   v                                  v
HelloController 后端接口控制器
   |
   v
UserRepository 数据访问层
   |
   v
User 实体类
   |
   v
MySQL 数据库 demo.user 表""")

add_heading(doc, "三、登录与注册流程", 1)
add_code_block(doc, """打开网站
   |
   v
index.html 自动跳转到 login.html
   |
   +-- 已有账号：输入账号和密码 -> POST /login
   |       |
   |       +-- 账号不存在：提示先注册
   |       +-- 密码错误：提示重新输入
   |       +-- 校验成功：显示欢迎信息
   |
   +-- 没有账号：进入 register.html
           |
           +-- 输入账号、密码、确认密码
           +-- 两次密码不一致：前端提示错误
           +-- 两次密码一致：POST /register
                   |
                   +-- 用户名已存在：注册失败
                   +-- 用户名不存在：保存到 MySQL，注册成功""")

add_heading(doc, "四、工程目录结构", 1)
add_table(
    doc,
    ["目录或文件", "作用"],
    [
        ("src/main/java/com/example/demo", "后端 Java 源码目录，包含启动类、控制器、实体类和数据库访问接口。"),
        ("src/main/resources", "应用配置和静态资源目录。"),
        ("src/main/resources/static", "前端页面目录，Spring Boot 会自动把这里的 HTML 作为静态资源提供给浏览器。"),
        ("src/test/java", "测试代码目录。"),
        ("pom.xml", "Maven 项目配置，管理依赖、插件、Java 版本和打包方式。"),
        ("target", "Maven 编译和打包后生成的目录，包括 class 文件和可运行 jar。"),
        (".idea", "IntelliJ IDEA 的本地项目配置。"),
        ("mvnw / mvnw.cmd", "Maven Wrapper 脚本，用于在没有全局 Maven 的情况下构建项目。"),
    ],
    [3000, 6360],
)

add_heading(doc, "五、核心源码解释", 1)

add_heading(doc, "1. DemoApplication.java", 2)
add_body(doc, "这是项目启动类。@SpringBootApplication 表示这是一个 Spring Boot 应用，会自动扫描 com.example.demo 包下的组件。main 方法调用 SpringApplication.run，启动内置 Tomcat、加载配置并运行整个 Web 应用。")
add_code_block(doc, """@SpringBootApplication
public class DemoApplication {
    public static void main(String[] args) {
        SpringApplication.run(DemoApplication.class, args);
    }
}""")

add_heading(doc, "2. HelloController.java", 2)
add_body(doc, "这是后端接口控制器，负责处理登录和注册请求。@RestController 表示方法返回值会直接转换成 JSON；@Autowired 自动注入 UserRepository，用来查询和保存用户。")
add_body(doc, "login 方法对应 POST /login，前端以表单方式提交 username 和 password。方法先校验空值，再根据用户名查数据库，然后判断用户是否存在、密码是否正确，最终返回 code、message 和 username。")
add_body(doc, "register 方法对应 POST /register，前端以 JSON 方式提交 username 和 password。方法先校验空值，再检查用户名是否重复，最后创建 User 对象并保存到数据库。")
add_table(
    doc,
    ["接口", "请求数据", "主要逻辑", "返回含义"],
    [
        ("/login", "表单参数 username、password", "空值校验；按用户名查询；判断密码", "code=0 成功；code=1 参数或密码错误；code=2 账号不存在"),
        ("/register", "JSON：username、password", "空值校验；检查重名；保存新用户", "success=true 注册成功；success=false 注册失败"),
    ],
    [1700, 2200, 3000, 2460],
)
add_code_block(doc, """User user = userRepository.findByUsername(username.trim());
if (user == null) {
    result.put("code", 2);
    result.put("message", "账号不存在，请先注册");
    return result;
}

if (!user.getPassword().equals(password)) {
    result.put("code", 1);
    result.put("message", "密码错误，请重新输入");
    return result;
}""")

add_heading(doc, "3. User.java", 2)
add_body(doc, "这是用户实体类，对应数据库中的用户表。@Entity 表示该类会被 JPA 映射为数据库表；@Id 表示主键；@GeneratedValue(strategy = GenerationType.IDENTITY) 表示主键由数据库自增生成。")
add_body(doc, "字段 id 保存用户编号，username 保存账号，password 保存密码。无参构造方法是 JPA 创建对象时需要的；带参构造方法用于注册时快速创建用户。getter/setter 用于读取和修改字段。toString 方法只输出 id 和 username，没有输出 password，避免日志中直接暴露密码。")
add_code_block(doc, """@Entity
public class User {
    @Id
    @GeneratedValue(strategy = GenerationType.IDENTITY)
    private Long id;
    private String username;
    private String password;
}""")

add_heading(doc, "4. UserRepository.java", 2)
add_body(doc, "这是数据库访问层接口。它继承 JpaRepository<User, Long> 后，自动拥有 save、findById、findAll、delete 等常用数据库操作方法。")
add_body(doc, "findByUsername 是 Spring Data JPA 的方法名查询。Spring 会根据方法名自动生成按 username 查询的 SQL，效果类似 select * from user where username = ?。")
add_code_block(doc, """@Repository
public interface UserRepository extends JpaRepository<User, Long> {
    User findByUsername(String username);
}""")

add_heading(doc, "六、配置文件解释", 1)
add_heading(doc, "application.properties", 2)
add_table(
    doc,
    ["配置项", "含义"],
    [
        ("spring.application.name=demo", "设置应用名称为 demo。"),
        ("spring.datasource.url=...", "配置 MySQL 连接地址，连接本机 3306 端口的 demo 数据库，并指定 UTF-8 和 Asia/Shanghai 时区。"),
        ("spring.datasource.username=root", "数据库用户名。"),
        ("spring.datasource.password=000906", "数据库密码。"),
        ("spring.datasource.driver-class-name=com.mysql.cj.jdbc.Driver", "指定 MySQL 8 驱动类。"),
        ("spring.jpa.hibernate.ddl-auto=update", "让 Hibernate 根据实体类自动更新数据库表结构。"),
        ("spring.jpa.show-sql=true", "在日志中打印 Hibernate 执行的 SQL。"),
        ("spring.jpa.properties.hibernate.dialect=org.hibernate.dialect.MySQL5Dialect", "指定 Hibernate 使用 MySQL 方言生成 SQL。"),
    ],
    [3600, 5760],
)

add_heading(doc, "七、前端页面代码解释", 1)
add_heading(doc, "1. index.html", 2)
add_body(doc, "这是首页，核心代码是 meta refresh。浏览器打开根路径后会立刻跳转到 login.html。如果自动跳转失败，页面中还提供了一个手动点击的登录页链接。")
add_code_block(doc, """<meta http-equiv="refresh" content="0;url=login.html">""")

add_heading(doc, "2. login.html", 2)
add_body(doc, "这是登录页面。HTML 部分包含账号输入框、密码输入框、登录按钮、错误提示区域、注册入口和登录成功后的欢迎区域。CSS 负责渐变背景、居中卡片、按钮、输入框和提示文字样式。")
add_body(doc, "JavaScript 中的 login 函数读取账号密码，先做空值校验，然后使用 fetch 发送 POST /login 请求。请求头 Content-Type 是 application/x-www-form-urlencoded，因此后端用 @RequestParam 接收参数。")
add_body(doc, "当后端返回 code=0 时，页面隐藏登录表单并显示欢迎信息；当 code=2 时，提示账号不存在并给出注册链接；其他情况显示密码错误或登录失败提示。logout 函数只是清空页面状态，没有真正的服务端会话退出逻辑。")
add_code_block(doc, """fetch('/login', {
    method: 'POST',
    headers: { 'Content-Type': 'application/x-www-form-urlencoded' },
    body: 'username=' + encodeURIComponent(username) +
          '&password=' + encodeURIComponent(password)
})""")

add_heading(doc, "3. register.html", 2)
add_body(doc, "这是注册页面。HTML 部分包含账号、密码、确认密码输入框、注册按钮、消息提示区域和返回登录页链接。CSS 负责注册卡片、输入框、按钮和成功/失败消息样式。")
add_body(doc, "JavaScript 监听表单 submit 事件并阻止默认刷新。它先读取三个输入值，检查两次密码是否一致；如果一致，就用 fetch 发送 POST /register 请求，请求体是 JSON，因此后端用 @RequestBody 接收。")
add_body(doc, "注册成功后，页面显示成功消息，并在 3 秒后跳转回 /login.html；注册失败时，直接显示后端返回的错误消息，例如用户名已存在。")
add_code_block(doc, """const res = await fetch('/register', {
    method: 'POST',
    headers: { 'Content-Type': 'application/json' },
    body: JSON.stringify({ username, password })
});""")

add_heading(doc, "八、测试与构建文件", 1)
add_table(
    doc,
    ["文件", "代码或内容说明"],
    [
        ("DemoApplicationTests.java", "@SpringBootTest 会启动 Spring 应用上下文。contextLoads 是默认空测试，用于验证项目配置是否能正常加载。"),
        ("pom.xml", "定义 Spring Boot 2.7.18、Java 1.8、Web、JPA、MySQL 驱动、测试依赖和 Spring Boot 打包插件。"),
        ("mvnw / mvnw.cmd", "Maven Wrapper 启动脚本，分别用于类 Unix 系统和 Windows。"),
        (".mvn/wrapper/maven-wrapper.properties", "指定 Maven Wrapper 使用 Maven 3.9.16。"),
        (".gitignore", "指定 Git 忽略 target、IDE 配置、构建目录等不需要提交的文件。"),
        (".gitattributes", "规定脚本文件换行符，例如 mvnw 使用 LF，cmd 使用 CRLF。"),
        ("HELP.md", "Spring Initializr 生成的项目帮助文档，包含 Maven 和 Spring Boot 参考链接。"),
        ("app.log", "应用运行日志，能看到项目曾在 8080 端口启动，并记录 Hibernate 查询和插入用户的 SQL。"),
    ],
    [2700, 6660],
)

add_heading(doc, "九、target 目录说明", 1)
add_body(doc, "target 目录是 Maven 编译和打包后的产物目录，不是主要源码。它可以删除，重新执行 Maven 构建后会再次生成。")
add_table(
    doc,
    ["target 中的文件", "说明"],
    [
        ("target/classes/application.properties", "从 src/main/resources 复制过来的运行配置。"),
        ("target/classes/static/*.html", "静态页面的构建输出副本。"),
        ("target/classes/com/example/demo/*.class", "Java 源码编译后的字节码文件。"),
        ("target/test-classes/*.class", "测试源码编译后的字节码文件。"),
        ("target/demo-0.0.1-SNAPSHOT.jar", "最终可运行的 Spring Boot Jar 包。"),
        ("target/demo-0.0.1-SNAPSHOT.jar.original", "Spring Boot 重新打包前的原始 Jar。"),
        ("target/maven-archiver/pom.properties", "记录 groupId、artifactId 和 version 等打包信息。"),
        ("target/maven-status/**/*.lst", "Maven 编译状态文件，记录参与编译的源码和生成的 class 文件。"),
    ],
    [3500, 5860],
)

add_heading(doc, "十、总结", 1)
add_body(doc, "这个项目代码量不大，但覆盖了一个 Web 登录注册系统的基本链路：前端页面收集输入，JavaScript 调用后端接口，Controller 处理请求，Repository 操作数据库，JPA 将 User 实体映射到 MySQL 表。")
add_body(doc, "如果继续完善，可以增加密码加密、登录状态保持、统一异常处理、表单规则校验、接口测试、前后端分离结构，以及更完整的用户信息管理功能。")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Examle1 工程说明文档")
    set_run_font(run, size=9, color="777777")

doc.save(OUT)
print(OUT)
